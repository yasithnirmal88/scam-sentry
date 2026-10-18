import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_revised_docx():
    doc = docx.Document()

    # Set page margins to standard 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    primary_color = RGBColor(15, 34, 64)      # Dark Navy
    secondary_color = RGBColor(41, 128, 185)  # Muted Blue
    text_dark = RGBColor(33, 37, 41)         # Off-black

    def set_cell_shading(cell, color_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def format_cell(cell, text, bold=False, font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, color=text_dark):
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        return run

    # 1. Document Title / Header
    h1 = doc.add_paragraph()
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(2)
    r1 = h1.add_run('ESSENTIALS OF ARTIFICIAL INTELLIGENCE')
    r1.font.name = 'Calibri'
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = primary_color

    h2 = doc.add_paragraph()
    h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h2.paragraph_format.space_before = Pt(0)
    h2.paragraph_format.space_after = Pt(12)
    r2 = h2.add_run('INDIVIDUAL CONTRIBUTION SHEET')
    r2.font.name = 'Calibri'
    r2.font.size = Pt(15)
    r2.font.bold = True
    r2.font.color.rgb = secondary_color

    # 2. Student & Project Metadata Table
    meta_table = doc.add_table(rows=3, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_widths = [Inches(1.5), Inches(2.2), Inches(1.5), Inches(1.8)]

    data_meta = [
        [('Student Name:', True), ('Yasith Nirmal Peiris', False), ('Registration No.:', True), ('D/BSE/25/0018', False)],
        [('Degree / Intake:', True), ('Information Technology / 2024', False), ('Group No.:', True), ('44', False)],
        [('Project Title:', True), ('Scam Sentry: An Intelligent Fake Job Advertisement Detection System', False), ('GitHub Username:', True), ('yasithnirmal88', False)]
    ]

    for r_idx, row_data in enumerate(data_meta):
        row = meta_table.rows[r_idx]
        for c_idx, cell_data in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            is_header = cell_data[1]
            set_cell_shading(cell, 'F8F9FA' if is_header else 'FFFFFF')
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            format_cell(cell, cell_data[0], bold=cell_data[1], font_size=9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3. Executive Contribution Summary
    sum_heading = doc.add_paragraph()
    sum_heading.paragraph_format.space_before = Pt(6)
    sum_heading.paragraph_format.space_after = Pt(2)
    sh_run = sum_heading.add_run('Executive Contribution Summary')
    sh_run.font.name = 'Calibri'
    sh_run.font.size = Pt(11)
    sh_run.font.bold = True
    sh_run.font.color.rgb = primary_color

    sum_p = doc.add_paragraph()
    sum_p.paragraph_format.space_before = Pt(0)
    sum_p.paragraph_format.space_after = Pt(8)
    sum_p.paragraph_format.line_spacing = 1.15
    sum_run = sum_p.add_run(
        'As the lead backend architect, risk engine designer, security engineer, and test author for Scam Sentry, '
        'I designed and developed the core intelligence and infrastructure of the platform (50 commits, >18,300 lines added across project). '
        'I engineered the deterministic multi-tier Risk Engine featuring 18 graded indicators (C1-L3), first-match-wins decision logic, '
        'and consistency safety nets to guarantee Critical threat dominance over probabilistic ML scores. I developed the pluggable 10-source '
        'job-board scanner, background daemon scheduler, and dual-backend storage facade supporting SQLite and PostgreSQL. Furthermore, I authored '
        'the complete 403-test automated quality assurance suite, established the CI/CD pipeline, hardened API security against brute-force and scraping abuse, '
        'and implemented the CV relevance matching and bare-city geographic resolution pipelines.'
    )
    sum_run.font.name = 'Calibri'
    sum_run.font.size = Pt(9.5)
    sum_run.font.color.rgb = text_dark

    # 4. Detailed Technical Contributions Table
    table_heading = doc.add_paragraph()
    table_heading.paragraph_format.space_before = Pt(4)
    table_heading.paragraph_format.space_after = Pt(2)
    th_run = table_heading.add_run('Detailed Technical Contributions')
    th_run.font.name = 'Calibri'
    th_run.font.size = Pt(11)
    th_run.font.bold = True
    th_run.font.color.rgb = primary_color

    tech_table = doc.add_table(rows=1, cols=5)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False

    tech_col_widths = [Inches(0.4), Inches(1.5), Inches(2.6), Inches(1.2), Inches(1.3)]

    headers = ['No.', 'Technical Area', 'My Specific Contribution', 'Tools & AI Techniques', 'Evidence & Deliverables']
    hdr_row = tech_table.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = tech_col_widths[idx]
        set_cell_shading(cell, '1F3A60') # Royal navy
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        format_cell(cell, heading, bold=True, font_size=9, align=WD_ALIGN_PARAGRAPH.CENTER if idx==0 else WD_ALIGN_PARAGRAPH.LEFT, color=RGBColor(255, 255, 255))

    contributions = [
        (
            '1',
            'Deterministic Risk Engine & Indicators',
            'Designed and implemented the core rule-based Risk Engine with 18 graded indicators across 4 severity tiers (C1-C4 Critical, H1-H6 High, M1-M5 Medium, L1-L3 Low). Implemented first-match-wins arbitration rules and V1-V3 consistency safety nets so Critical fraud indicators (advance fee, sensitive identity requests) can never be averaged away by ML probability scores.',
            'Rule-based Reasoning, Python, Regex Engine, Deterministic Decision Trees',
            'risk/indicators.py, risk/engine.py, risk/consistency.py (Commits fe0abfd, 4cec4a1, 0c191b8; tests/test_engine.py 17 passed)'
        ),
        (
            '2',
            'Responsible AI Language Discipline',
            'Engineered ethical language discipline sanitization preventing defamatory allegations (e.g., claiming a listing is "fake" or "a scam"). Generated forensic, evidence-backed red flag statements coupled with legitimate alternative business explanations, ensuring academic and legal safety.',
            'Deterministic Lexical Sweep, Content Moderation, NLP Sanitization',
            'risk/language.py (Commit 4cec4a1, 0c191b8; tests/test_language.py 5 passed)'
        ),
        (
            '3',
            'Job Board Scanner Framework & Ingestion',
            'Built an extensible, pluggable job ingestion architecture (BaseScanner interface) and implemented connectors for 10 public job boards (RemoteOK, Remotive, Arbeitnow, Jobicy, Himalayas, Working Nomads, LaraJobs, Greenhouse, Lever, Ashby). Built exponential-backoff HTTP retry logic and SHA-256 URL hashing for deduplication.',
            'Python Requests, REST/JSON & RSS APIs, Regex HTML Stripping, SHA-256 Deduplication',
            'scanner/base.py, scanner/registry.py, scanner/sources/*.py (Commits 4338d3e, bb965f9, 5d6d4cb; 24 scanner tests passed)'
        ),
        (
            '4',
            'Background Scan Orchestration & Scheduler',
            'Developed an asynchronous background scanner service utilizing Python daemon worker threads with thread-safe execution locks, preventing concurrent duplicate scans. Added scan history logging, CLI scanning tool (scan.py), and non-blocking HTTP 202 endpoints with synchronous bypass mode.',
            'Python Threading, Daemon Scheduling, CLI Engineering, Mutex Concurrency',
            'services/scheduler.py, services/scanner_service.py, scan.py (Commits bb965f9, b4fbaf5; tests/test_phase3.py 20 passed)'
        ),
        (
            '5',
            'Dual-Engine Storage Abstraction & Migration',
            'Designed a unified repository facade (models/job.py, storage.py) supporting SQLite for local zero-config development and PostgreSQL for production via DATABASE_URL. Built automatic schema evolution, dynamic index creation, and idempotent background backfilling for historical records.',
            'SQLite, PostgreSQL, psycopg3, Database Migration, Repository Pattern',
            'models/job.py, models/storage.py (Commits 4338d3e, 4e074a4, 0e09b6b; tests/test_job_model.py 17 passed)'
        ),
        (
            '6',
            'API Security, Rate Limiting & Abuse Control',
            'Engineered enterprise security remediations (Audit F-01 to F-08): fail-closed X-Scan-Token authentication, per-IP sliding window rate limiting (Flask-Limiter), strict Content Security Policy (CSP), anti-clickjacking headers (SAMEORIGIN), payload size limits (JSON 413), and fail-closed application boot guards.',
            'Flask-Limiter, HTTP Security Headers, API Token Auth, Abuse Mitigation',
            'app.py, config.py (Commits 1ec4d9c, 7ed03f4; tests/test_abuse_control.py 23 passed, test_scan_api.py 22 passed)'
        ),
        (
            '7',
            'ML Calibration & Supporting Signal Honesty',
            'Retrained GradientBoosting and TF-IDF models to resolve library depreciation warnings, calibrated model probability curves via isotonic regression, and transparently labeled ML scores as synthetic-trained supporting signals that cannot override deterministic risk floors.',
            'Scikit-learn, TF-IDF Vectorization, Gradient Boosting, Isotonic Calibration',
            'model_training.py, artifacts/, app.py (Commits 0969ed9, dfbc1db; tests/test_ml_honesty.py 14 passed)'
        ),
        (
            '8',
            'Domain Taxonomy & Bare-City Geo Normalization',
            'Created an 18-domain, 100+ role keyword taxonomy with boundary-aware regex patterns. Built a geographic normalizer parsing continents, countries, and an 80+ bare-city lookup table (e.g. "London" -> UK, "Bangalore" -> India) with automated database migration backfilling.',
            'Regex Pattern Matching, Lexical Mapping, Entity Resolution, Geographic Normalization',
            'taxonomy.py, location.py (Commits 34238f1, fba3fd1, 2366f41, 90e2e8e; 25 location/taxonomy tests passed)'
        ),
        (
            '9',
            'Resume / CV Career Relevance Matcher',
            'Developed a candidate CV relevance matching engine utilizing TF-IDF n-gram vectorization and cosine similarity to rank ingested postings against candidate resumes. Explicitly isolated CV matching from risk scoring to prevent career relevance from biasing fraud detection.',
            'NLP, TF-IDF Cosine Similarity, Resume Feature Extraction, Vector Space Model',
            'cv_matcher.py, cv_utils.py (Commits cf8f90a, 65e6473; tests/test_cv_matcher.py 10 passed, test_cv_api.py 8 passed)'
        ),
        (
            '10',
            'Automated Testing, CI/CD & Deployment',
            'Engineered the comprehensive 403-test automated quality assurance suite (unit, integration, hostile fuzzing, and security tests). Built the GitHub Actions CI workflow spinning up ephemeral PostgreSQL containers for storage parity and configured Docker runtime packaging for Render cloud deployment.',
            'Pytest, GitHub Actions, Docker, Render Cloud PaaS, Continuous Integration',
            'fakejob_app/tests/*, .github/workflows/ci.yml, Dockerfile (Commits fe0abfd, 0c191b8, f5a7156; 403 tests passed in 18.89s)'
        )
    ]

    for row_data in contributions:
        row = tech_table.add_row()
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        for idx, text in enumerate(row_data):
            cell = row.cells[idx]
            cell.width = tech_col_widths[idx]
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            bold = True if idx in [0, 1] else False
            font_size = 8.5
            format_cell(cell, text, bold=bold, font_size=font_size, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5. Tools & Technologies Summary
    tech_hdr = doc.add_paragraph()
    tech_hdr.paragraph_format.space_before = Pt(6)
    tech_hdr.paragraph_format.space_after = Pt(2)
    th2_run = tech_hdr.add_run('Summary of Tools, Frameworks & AI Techniques Used')
    th2_run.font.name = 'Calibri'
    th2_run.font.size = Pt(11)
    th2_run.font.bold = True
    th2_run.font.color.rgb = primary_color

    tech_p = doc.add_paragraph()
    tech_p.paragraph_format.space_before = Pt(0)
    tech_p.paragraph_format.space_after = Pt(8)
    t_run = tech_p.add_run(
        '• Programming & Frameworks: Python 3.11+, Flask 3.0, Flask-Limiter, Jinja2, Requests, Threading.\n'
        '• Artificial Intelligence & NLP: Scikit-learn (GradientBoosting, TF-IDF Vectorizer, Cosine Similarity), Rule-based Expert Systems, Probability Calibration, Deterministic Regex Decision Trees.\n'
        '• Database & Persistence: SQLite3, PostgreSQL (psycopg3), Database Schema Migrations, Repository Design Pattern.\n'
        '• Quality Assurance & Security: Pytest (403 tests), GitHub Actions CI/CD, Pyflakes, API Token Authentication, CSP/OWASP Security Headers.\n'
        '• Deployment & Infrastructure: Docker, Render Cloud PaaS, Tesseract OCR, Poppler Utilities, pdfplumber.'
    )
    t_run.font.name = 'Calibri'
    t_run.font.size = Pt(9)
    t_run.font.color.rgb = text_dark

    # 6. Student Declaration Section
    dec_hdr = doc.add_paragraph()
    dec_hdr.paragraph_format.space_before = Pt(6)
    dec_hdr.paragraph_format.space_after = Pt(2)
    dh_run = dec_hdr.add_run('Declaration')
    dh_run.font.name = 'Calibri'
    dh_run.font.size = Pt(11)
    dh_run.font.bold = True
    dh_run.font.color.rgb = primary_color

    dec_p = doc.add_paragraph()
    dec_p.paragraph_format.space_before = Pt(0)
    dec_p.paragraph_format.space_after = Pt(6)
    dec_p.paragraph_format.line_spacing = 1.15
    d_run = dec_p.add_run(
        'I confirm that the technical contributions, architectural designs, implementations, and test suites detailed above '
        'accurately and truthfully represent my individual work in Group 44 for the Scam Sentry project. I am prepared to demonstrate, '
        'defend, and explain all associated source code, algorithms, and technical artifacts during oral examination.'
    )
    d_run.font.name = 'Calibri'
    d_run.font.size = Pt(9.5)
    d_run.font.italic = True
    d_run.font.color.rgb = text_dark

    # Signature Table
    sig_table = doc.add_table(rows=1, cols=4)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    sig_col_widths = [Inches(1.5), Inches(2.2), Inches(1.5), Inches(1.8)]
    sig_row = sig_table.rows[0]

    sig_cells_data = [
        ('Student Signature:', True),
        ('_____________________________', False),
        ('Submission Date:', True),
        ('___ / ___ / 2026', False)
    ]

    for idx, cell_data in enumerate(sig_cells_data):
        cell = sig_row.cells[idx]
        cell.width = sig_col_widths[idx]
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        set_cell_shading(cell, 'F8F9FA' if cell_data[1] else 'FFFFFF')
        format_cell(cell, cell_data[0], bold=cell_data[1], font_size=9.5)

    output_path = r'd:\GitHub\Fake_job_detector\Individual_Contribution_Sheet_Yasith_Nirmal_Peiris_REVISED.docx'
    doc.save(output_path)
    print(f'Successfully generated {output_path}')

if __name__ == '__main__':
    create_revised_docx()
